"""Gera a apresentação do Learning Helper em formato PPTX.

O arquivo é compatível com o Google Slides (Drive → Novo → Upload de arquivo).

Uso:
    uv run --with python-pptx python scripts/build_presentation.py
    uv run --with python-pptx python scripts/build_presentation.py --out /tmp/deck.pptx
"""

from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
DEFAULT_OUT = ROOT / "docs" / "apresentacao" / "learning-helper.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

FONT = "Arial"

NAVY = "0F172A"
NAVY_SOFT = "17233D"
INK = "111827"
MUTED = "64748B"
MUTED_DARK = "94A3B8"
WHITE = "FFFFFF"
LIGHT = "F8FAFC"
BORDER = "E2E8F0"
SHADOW = "E9EEF5"

BLUE = "2563EB"
BLUE_SOFT = "EFF6FF"
TEAL = "0D9488"
TEAL_SOFT = "F0FDFA"
AMBER = "D97706"
AMBER_SOFT = "FFF7ED"
GREEN = "15803D"
GREEN_SOFT = "F0FDF4"
RED = "C2410C"
RED_SOFT = "FFF7ED"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value.replace("#", "").upper())


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 17,
    bold: bool = False,
    italic: bool = False,
    value: str = INK,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing: float = 1.0,
    link: str | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(value)
    if link:
        run.hyperlink.address = link
    return box


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    border: str | None = BORDER,
    radius: float = 0.08,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color(border)
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    try:
        shape.adjustments[0] = radius
    except (IndexError, ValueError):
        pass
    return shape


def add_rect(slide, x: float, y: float, w: float, h: float, fill: str):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_oval(slide, x: float, y: float, w: float, h: float, fill: str):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_picture(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    border: str | None = None,
    alt: str = "",
):
    with Image.open(path) as image:
        image_w, image_h = image.size
    scale = min(w / image_w, h / image_h)
    picture_w = image_w * scale
    picture_h = image_h * scale
    picture_x = x + (w - picture_w) / 2
    picture_y = y + (h - picture_h) / 2
    picture = slide.shapes.add_picture(
        str(path),
        Inches(picture_x),
        Inches(picture_y),
        Inches(picture_w),
        Inches(picture_h),
    )
    if border:
        picture.line.color.rgb = color(border)
        picture.line.width = Pt(1)
    else:
        picture.line.fill.background()
    picture._element.nvPicPr.cNvPr.set("descr", alt)
    return picture


def add_framed_image(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    padding: float = 0.1,
    alt: str = "",
):
    add_card(slide, x + 0.05, y + 0.06, w, h, fill=SHADOW, border=None, radius=0.04)
    add_card(slide, x, y, w, h, fill=WHITE, border=BORDER, radius=0.04)
    add_picture(
        slide,
        path,
        x + padding,
        y + padding,
        w - 2 * padding,
        h - 2 * padding,
        alt=alt,
    )


def set_title(
    slide,
    text: str,
    *,
    x: float = 0.7,
    y: float = 0.82,
    w: float = 11.9,
    h: float = 0.72,
    size: float = 30,
    value: str = INK,
):
    title = slide.shapes.title
    if title is None:
        return add_text(slide, text, x, y, w, h, size=size, bold=True, value=value)
    title.left = Inches(x)
    title.top = Inches(y)
    title.width = Inches(w)
    title.height = Inches(h)
    frame = title.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 0.98
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = color(value)


def add_background(slide, fill: str):
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = color(fill)


def add_footer(slide, number: int, *, dark: bool = False):
    line_color = "334155" if dark else BORDER
    text_color = MUTED_DARK if dark else MUTED
    add_rect(slide, 0.7, 7.08, 11.93, 0.01, line_color)
    add_text(slide, "Learning Helper", 0.7, 7.14, 5.0, 0.2, size=9, value=text_color)
    add_text(
        slide,
        str(number),
        11.9,
        7.14,
        0.73,
        0.2,
        size=9,
        value=text_color,
        align=PP_ALIGN.RIGHT,
    )


def add_header(slide, number: int, kicker: str, title: str, accent: str):
    add_rect(slide, 0, 0, SLIDE_W, 0.1, accent)
    add_text(
        slide,
        kicker.upper(),
        0.7,
        0.48,
        8.0,
        0.25,
        size=10,
        bold=True,
        value=accent,
    )
    set_title(slide, title)
    add_footer(slide, number)


def add_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def panel_crop(source: Path, target: Path) -> Path:
    with Image.open(source) as image:
        image = image.convert("RGB")
        width, height = image.size
        box = (
            int(width * 0.17),
            int(height * 0.18),
            int(width * 0.83),
            int(height * 0.82),
        )
        image.crop(box).save(target, optimize=True)
    return target


def stats_crop(source: Path, target: Path) -> Path:
    with Image.open(source) as image:
        image = image.convert("RGB")
        width, height = image.size
        image.crop((int(width * 0.19), 0, width, height)).save(target, optimize=True)
    return target


def build(out: Path, assets: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    title_layout = prs.slide_layouts[5]
    blank_layout = prs.slide_layouts[6]

    challenge = {
        number: panel_crop(
            SHOTS / f"{number:02d}-desafio-{slug}.png",
            assets / f"challenge-{number}.png",
        )
        for number, slug in (
            (1, "cores"),
            (2, "cores-diferentes"),
            (3, "comparacao"),
            (4, "contas"),
            (5, "problemas"),
            (6, "perguntas"),
            (7, "leitura"),
        )
    }
    bathroom = panel_crop(SHOTS / "10-lembrete-banheiro.png", assets / "bathroom.png")
    settings = SHOTS / "13-configuracoes-topo.png"
    stats_evolution = stats_crop(
        SHOTS / "16-estatisticas-evolucao.png", assets / "stats-evolution.png"
    )

    # ---------------------------------------------------------------- capa
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, NAVY)
    add_oval(slide, 10.4, -1.8, 4.8, 4.8, NAVY_SOFT)
    add_oval(slide, -2.6, 5.2, 4.2, 4.2, "0B2E36")

    add_text(
        slide,
        "APRESENTAÇÃO DO PROJETO",
        0.8,
        0.82,
        7.0,
        0.3,
        size=10.5,
        bold=True,
        value="5EEAD4",
    )
    title = slide.shapes.title
    title.left = Inches(0.8)
    title.top = Inches(1.25)
    title.width = Inches(6.3)
    title.height = Inches(1.65)
    frame = title.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(("Learning", "Helper")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 0.92
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(52)
        run.font.bold = True
        run.font.color.rgb = color(WHITE)

    add_text(
        slide,
        "Aprendizagem e rotina no ritmo certo.",
        0.85,
        3.05,
        6.2,
        0.55,
        size=22,
        bold=True,
        value=WHITE,
    )
    add_text(
        slide,
        "Desafios curtos, previsíveis e adaptativos para apoiar o "
        "desenvolvimento cognitivo e a autonomia de uma criança autista.",
        0.85,
        3.8,
        5.95,
        1.25,
        size=15,
        value="CBD5E1",
        line_spacing=1.2,
    )
    link = add_text(
        slide,
        "github.com/herossa/learning_helper  ↗",
        0.85,
        6.1,
        5.7,
        0.3,
        size=11,
        bold=True,
        value="5EEAD4",
    )
    link.click_action.hyperlink.address = "https://github.com/herossa/learning_helper"
    add_text(
        slide,
        "Uma visão simples do que é, por que existe e como funciona.",
        0.85,
        6.5,
        5.8,
        0.3,
        size=10,
        value=MUTED_DARK,
    )

    add_framed_image(
        slide,
        challenge[1],
        7.2,
        1.15,
        5.55,
        3.9,
        padding=0.12,
        alt="Tela do desafio de cores",
    )
    add_text(
        slide,
        "Desafio de cores  •  nível 1",
        7.35,
        5.13,
        5.25,
        0.25,
        size=9.5,
        value=MUTED_DARK,
    )
    for index, label in enumerate(("7 módulos", "5 níveis", "1–5 min")):
        x = 7.2 + index * 1.85
        add_card(slide, x, 5.65, 1.7, 0.4, fill=NAVY_SOFT, border="334155", radius=0.5)
        add_text(
            slide,
            label,
            x,
            5.73,
            1.7,
            0.22,
            size=9.5,
            bold=True,
            value="E2E8F0",
            align=PP_ALIGN.CENTER,
        )
    add_notes(
        slide,
        "Esta é uma visão geral do Learning Helper. O projeto nasceu para "
        "apoiar a aprendizagem e a rotina do meu filho autista, sem "
        "transformar o dia a dia em uma sequência de cobranças.",
    )

    # ------------------------------------------------------------ motivação
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 2, "Motivação", "Nasceu de uma necessidade real", AMBER)

    cards = (
        (
            BLUE,
            "Desenvolver habilidades",
            "Cognição, raciocínio lógico, linguagem e atenção em pequenas doses.",
        ),
        (
            TEAL,
            "Apoiar a rotina",
            "Lembretes de banheiro nos intervalos certos, com transição clara.",
        ),
        (
            AMBER,
            "Aprender sem sobrecarga",
            "Layout estável, instruções curtas e feedback gentil.",
        ),
    )
    for index, (accent, heading, body) in enumerate(cards):
        y = 1.65 + index * 1.5
        add_card(slide, 0.7, y, 6.15, 1.33, fill=WHITE, border=BORDER)
        add_rect(slide, 0.7, y, 0.09, 1.33, accent)
        add_text(slide, heading, 1.0, y + 0.18, 5.55, 0.3, size=14, bold=True)
        add_text(
            slide,
            body,
            1.0,
            y + 0.55,
            5.45,
            0.55,
            size=11.5,
            value=MUTED,
            line_spacing=1.15,
        )

    add_framed_image(
        slide,
        bathroom,
        7.35,
        1.65,
        5.35,
        3.1,
        padding=0.1,
        alt="Lembrete de banheiro",
    )
    add_text(
        slide,
        "Lembrete de banheiro: parte do cuidado, não uma punição.",
        7.4,
        4.95,
        5.2,
        0.5,
        size=10.5,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Não substitui acompanhamento pedagógico, fonoaudiológico ou terapêutico.",
        0.7,
        6.45,
        11.9,
        0.3,
        size=10,
        value=MUTED,
    )
    add_notes(
        slide,
        "A motivação veio de uma necessidade concreta da nossa família: "
        "oferecer ao meu filho uma forma previsível de exercitar habilidades "
        "cognitivas e, ao mesmo tempo, apoiá-lo na rotina de ir ao banheiro. "
        "O app é um apoio, não um substituto de acompanhamento profissional.",
    )

    # ---------------------------------------------------------- como funciona
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 3, "Como funciona", "Um ciclo curto, simples e previsível", BLUE)
    add_text(
        slide,
        "O app aparece na hora certa, faz uma pergunta e devolve a tela.",
        0.7,
        1.6,
        11.9,
        0.35,
        size=13.5,
        value=MUTED,
    )

    steps = (
        ("1", "Agenda", "Intervalo de 1 a 5 min", BLUE, BLUE_SOFT),
        ("2", "Desafio", "Cobre a tela", TEAL, TEAL_SOFT),
        ("3", "Resposta", "Clique ou digite", AMBER, AMBER_SOFT),
        ("4", "Feedback", "Certo! / Quase!", BLUE, BLUE_SOFT),
        ("5", "Registro", "Tentativa no SQLite", TEAL, TEAL_SOFT),
        ("6", "Ajuste", "Sobe ou desce", AMBER, AMBER_SOFT),
    )
    card_w = 1.75
    gap = 0.286
    for index, (number, heading, body, accent, soft) in enumerate(steps):
        x = 0.7 + index * (card_w + gap)
        add_card(slide, x, 2.2, card_w, 2.2, fill=WHITE, border=BORDER)
        add_oval(slide, x + 0.55, 2.4, 0.65, 0.65, soft)
        add_text(
            slide,
            number,
            x + 0.55,
            2.54,
            0.65,
            0.3,
            size=15,
            bold=True,
            value=accent,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            heading,
            x + 0.12,
            3.22,
            card_w - 0.24,
            0.3,
            size=13,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            body,
            x + 0.12,
            3.62,
            card_w - 0.24,
            0.55,
            size=9.5,
            value=MUTED,
            align=PP_ALIGN.CENTER,
            line_spacing=1.05,
        )
        if index < len(steps) - 1:
            add_text(
                slide,
                "→",
                x + card_w + 0.035,
                2.95,
                0.22,
                0.4,
                size=18,
                bold=True,
                value=accent,
                align=PP_ALIGN.CENTER,
            )

    add_card(slide, 0.7, 4.95, 11.93, 1.35, fill=BLUE_SOFT, border=None)
    add_text(
        slide,
        "Intervalo inicial de 3 min  •  acerto +30 s  •  erro −30 s",
        1.0,
        5.2,
        11.3,
        0.35,
        size=17,
        bold=True,
        value=BLUE,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Os limites de 1 a 5 minutos e a dificuldade acompanham o ritmo da criança.",
        1.0,
        5.65,
        11.3,
        0.3,
        size=11,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "A tela é liberada logo depois do feedback.",
        0.7,
        6.55,
        11.9,
        0.3,
        size=10,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "O ciclo é curto de propósito. O app agenda o próximo desafio, cobre a "
        "tela por alguns instantes, registra a resposta e usa o resultado para "
        "ajustar o intervalo e a dificuldade. A criança não fica presa em uma "
        "sessão longa.",
    )

    # --------------------------------------------------------------- módulos
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 4, "Os módulos", "Sete tipos de desafio, uma coisa de cada vez", TEAL)
    add_text(
        slide,
        "Começa com Cores. Os demais são desbloqueados conforme o desempenho.",
        0.7,
        1.6,
        11.9,
        0.3,
        size=13,
        value=MUTED,
    )

    modules = (
        (1, "Cores"),
        (2, "Cores diferentes"),
        (3, "Comparação"),
        (4, "Contas"),
        (5, "Problemas"),
        (6, "Perguntas"),
        (7, "Leitura"),
    )
    card_w = 2.85
    card_h = 2.05
    gap_x = 0.2
    for index, (number, label) in enumerate(modules):
        row = index // 4
        column = index % 4
        x = 0.665 + column * (card_w + gap_x)
        y = 2.05 + row * (card_h + 0.18)
        add_card(slide, x, y, card_w, card_h, fill=WHITE, border=BORDER)
        add_picture(
            slide,
            challenge[number],
            x + 0.1,
            y + 0.1,
            card_w - 0.2,
            1.47,
            border="CBD5E1",
            alt=f"Desafio de {label.lower()}",
        )
        add_text(
            slide,
            f"{number:02d}  ·  {label}",
            x + 0.12,
            y + 1.65,
            card_w - 0.24,
            0.25,
            size=11,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    info_x = 0.665 + 3 * (card_w + gap_x)
    info_y = 2.05 + (card_h + 0.18)
    add_card(slide, info_x, info_y, card_w, card_h, fill=TEAL_SOFT, border="99F6E4")
    add_text(
        slide,
        "5 níveis",
        info_x + 0.15,
        info_y + 0.28,
        card_w - 0.3,
        0.35,
        size=18,
        bold=True,
        value=TEAL,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "por módulo",
        info_x + 0.15,
        info_y + 0.66,
        card_w - 0.3,
        0.25,
        size=10,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Cada módulo aparece uma vez antes de repetir.",
        info_x + 0.18,
        info_y + 1.16,
        card_w - 0.36,
        0.55,
        size=9.5,
        value=TEAL,
        align=PP_ALIGN.CENTER,
        line_spacing=1.1,
    )
    add_text(
        slide,
        "Módulos desativados são pulados, mas o progresso fica guardado.",
        0.7,
        6.55,
        11.9,
        0.3,
        size=10,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "São sete módulos: cores, cores diferentes, comparação, contas, "
        "problemas, perguntas e leitura. A criança começa com Cores, no nível 1, "
        "e os outros são liberados aos poucos. Cada módulo aparece uma vez "
        "antes de qualquer um repetir.",
    )

    # ------------------------------------------------------------- adaptação
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 5, "Adaptação", "O desafio acompanha o ritmo dele", AMBER)

    stats = (
        ("1–5", "níveis por módulo", BLUE, BLUE_SOFT),
        ("10", "tentativas por avaliação", TEAL, TEAL_SOFT),
        ("≥ 80%", "sobe de nível e desbloqueia", GREEN, GREEN_SOFT),
        ("≤ 40%", "desce de nível", RED, RED_SOFT),
    )
    for index, (number, label, accent, soft) in enumerate(stats):
        row = index // 2
        column = index % 2
        x = 0.7 + column * 2.75
        y = 1.6 + row * 1.5
        add_card(slide, x, y, 2.55, 1.3, fill=WHITE, border=BORDER)
        add_rect(slide, x, y, 0.07, 1.3, accent)
        add_text(slide, number, x + 0.25, y + 0.14, 2.1, 0.5, size=24, bold=True, value=accent)
        add_text(slide, label, x + 0.25, y + 0.72, 2.1, 0.35, size=10.5, value=MUTED)

    add_text(
        slide,
        "A cada bloco de 10 tentativas, o app calcula a precisão. O nível sobe, "
        "desce ou permanece — e o próximo módulo pode ser liberado.",
        0.7,
        4.85,
        5.35,
        0.9,
        size=11.5,
        value=INK,
        line_spacing=1.2,
    )
    add_text(
        slide,
        "Módulo já desbloqueado não desaparece se o desempenho cair.",
        0.7,
        6.35,
        5.35,
        0.35,
        size=9.5,
        value=MUTED,
    )

    add_text(
        slide,
        "Trilha de níveis",
        6.45,
        1.62,
        5.9,
        0.25,
        size=10.5,
        bold=True,
        value=MUTED,
    )
    add_text(
        slide,
        "sobe com ≥ 80%   →",
        6.45,
        1.98,
        5.9,
        0.3,
        size=10.5,
        bold=True,
        value=GREEN,
        align=PP_ALIGN.CENTER,
    )
    level_w = 1.05
    level_gap = 0.17
    for index in range(5):
        x = 6.45 + index * (level_w + level_gap)
        add_card(slide, x, 2.45, level_w, 0.95, fill=WHITE, border=BORDER)
        add_text(
            slide,
            str(index + 1),
            x,
            2.57,
            level_w,
            0.4,
            size=20,
            bold=True,
            value=BLUE if index == 2 else INK,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            "nível",
            x,
            3.04,
            level_w,
            0.2,
            size=8,
            value=MUTED,
            align=PP_ALIGN.CENTER,
        )
    add_text(
        slide,
        "←   desce com ≤ 40%",
        6.45,
        3.6,
        5.9,
        0.3,
        size=10.5,
        bold=True,
        value=RED,
        align=PP_ALIGN.CENTER,
    )
    add_card(slide, 6.45, 4.35, 5.9, 1.0, fill=BLUE_SOFT, border=None)
    add_text(
        slide,
        "O módulo seguinte é liberado quando o módulo mais novo atinge 80%.",
        6.65,
        4.38,
        5.5,
        0.94,
        size=11.5,
        bold=True,
        value=BLUE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        line_spacing=1.15,
    )
    add_text(
        slide,
        "Cada módulo tem sua própria trilha de 1 a 5.",
        6.45,
        5.6,
        5.9,
        0.3,
        size=9.5,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "A adaptação não olha uma resposta isolada: ela usa blocos de 10 "
        "tentativas. Com 80% ou mais de acerto, o nível sobe; com 40% ou menos, "
        "desce. O módulo seguinte só é liberado quando o mais novo atinge o "
        "critério. E um módulo já aberto não é fechado de novo.",
    )

    # --------------------------------------------------------------- banheiro
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 6, "Rotina", "Um lembrete que ajuda na transição", TEAL)

    add_framed_image(
        slide,
        bathroom,
        0.55,
        1.55,
        6.45,
        3.75,
        padding=0.1,
        alt="Tela do lembrete de banheiro",
    )
    add_text(
        slide,
        "Tela de lembrete com contagem regressiva.",
        0.7,
        5.45,
        6.1,
        0.3,
        size=10,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )

    routines = (
        "A cada 30 minutos, por padrão",
        "1 minuto de tela com contagem regressiva",
        "Se houver desafio na tela, espera e tenta de novo",
        "Independe do desempenho nos desafios",
    )
    for index, text in enumerate(routines):
        y = 1.7 + index * 0.82
        add_oval(slide, 7.25, y + 0.08, 0.28, 0.28, TEAL_SOFT)
        add_oval(slide, 7.33, y + 0.16, 0.12, 0.12, TEAL)
        add_text(slide, text, 7.7, y, 5.0, 0.55, size=12.5, value=INK, line_spacing=1.05)

    add_card(slide, 7.15, 5.15, 5.5, 0.95, fill=TEAL_SOFT, border=None)
    add_text(
        slide,
        "Faz parte do cuidado com a rotina — não é punição.",
        7.35,
        5.35,
        5.1,
        0.5,
        size=13,
        bold=True,
        value=TEAL,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "O intervalo e a duração podem ser ajustados nas configurações.",
        7.15,
        6.35,
        5.5,
        0.3,
        size=9.5,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "O lembrete de banheiro é independente dos desafios. Por padrão, ele "
        "aparece a cada 30 minutos e bloqueia a tela por 1 minuto, com uma "
        "contagem regressiva. Se um desafio estiver na tela, ele espera e tenta "
        "de novo. A ideia é apoiar a transição e a autonomia, não interromper "
        "o aprendizado como punição.",
    )

    # ---------------------------------------------------------- configurações
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 7, "Configurações", "Os adultos ajustam o que faz sentido", BLUE)

    add_framed_image(
        slide,
        settings,
        0.55,
        1.55,
        4.55,
        4.65,
        padding=0.08,
        alt="Janela de configurações",
    )
    add_text(
        slide,
        "A primeira execução já abre essa tela.",
        0.6,
        6.35,
        4.45,
        0.3,
        size=9.5,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )

    settings_items = (
        ("Intervalo", "mínimo, máximo, inicial e passo por resposta"),
        ("Avaliação", "janela de tentativas, limiares e nível máximo"),
        ("Ritmo", "duração do feedback e pausa do adulto"),
        ("Rotina", "lembrete de banheiro e módulos ativos"),
        ("Dados", "caminho do banco e verificação de atualizações"),
    )
    for index, (heading, body) in enumerate(settings_items):
        y = 1.65 + index * 0.88
        add_oval(slide, 5.4, y + 0.02, 0.36, 0.36, BLUE_SOFT)
        add_text(
            slide,
            str(index + 1),
            5.4,
            y + 0.08,
            0.36,
            0.22,
            size=10,
            bold=True,
            value=BLUE,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, heading, 5.95, y, 1.4, 0.28, size=12.5, bold=True)
        add_text(slide, body, 7.15, y, 5.25, 0.5, size=11, value=MUTED, line_spacing=1.1)

    add_card(slide, 5.4, 6.0, 7.0, 0.5, fill=BLUE_SOFT, border=None)
    add_text(
        slide,
        "As mudanças valem no próximo início do app.",
        5.6,
        6.12,
        6.6,
        0.25,
        size=10.5,
        bold=True,
        value=BLUE,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "Os adultos podem ajustar intervalo, critérios de avaliação, feedback, "
        "pausa, lembrete de banheiro, módulos ativos e atualizações. A tela "
        "aparece automaticamente na primeira execução. Na maior parte dos "
        "casos, as mudanças passam a valer no próximo início do app.",
    )

    # --------------------------------------------------------------- dados
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 8, "Acompanhamento", "Dados para acompanhar, não para pressionar", AMBER)

    add_framed_image(
        slide,
        stats_evolution,
        0.55,
        1.5,
        6.0,
        5.0,
        padding=0.08,
        alt="Painel de evolução",
    )
    add_text(
        slide,
        "Painel de evolução (dados de demonstração).",
        0.7,
        6.55,
        5.7,
        0.25,
        size=9,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )

    tabs = (
        ("Evolução", "Tentativas, acertos, precisão, dias ativos e gráficos por dia."),
        ("Módulos e níveis", "Nível atual, progresso e módulos desbloqueados."),
        ("Tentativas recentes", "Pergunta, resposta, gabarito e resultado de cada tentativa."),
    )
    for index, (heading, body) in enumerate(tabs):
        y = 1.7 + index * 1.35
        add_card(slide, 6.95, y, 5.65, 1.12, fill=WHITE, border=BORDER)
        add_rect(slide, 6.95, y, 0.07, 1.12, AMBER)
        add_text(slide, heading, 7.2, y + 0.16, 5.15, 0.28, size=13, bold=True)
        add_text(
            slide,
            body,
            7.2,
            y + 0.53,
            5.15,
            0.5,
            size=10.5,
            value=MUTED,
            line_spacing=1.1,
        )

    add_card(slide, 6.95, 5.85, 5.65, 0.65, fill=AMBER_SOFT, border=None)
    add_text(
        slide,
        "Somente leitura: não altera níveis, banco ou desafios.",
        7.15,
        6.03,
        5.25,
        0.3,
        size=10.5,
        bold=True,
        value=AMBER,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "As estatísticas ficam em um painel local, aberto no navegador. A aba "
        "de evolução mostra tentativas, acertos, precisão e dias ativos; a de "
        "módulos mostra níveis e progresso; a de tentativas recentes mostra "
        "pergunta, resposta e gabarito. O painel é somente leitura: ele não "
        "muda o banco, os níveis ou os desafios.",
    )

    # ---------------------------------------------------------- privacidade
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, LIGHT)
    add_header(slide, 9, "Tecnologia", "Simples de instalar. Local por padrão.", BLUE)

    tech = (
        ("01", "Dados locais", "SQLite na pasta do usuário. Nada é enviado para servidores."),
        ("02", "Instalação por usuário", "Windows e Linux, sem root por padrão."),
        ("03", "Atualização segura", "GitHub Releases com verificação SHA-256."),
        ("04", "App desktop", "Python + customtkinter, rodando junto com a rotina."),
    )
    for index, (number, heading, body) in enumerate(tech):
        x = 0.7 + index * 3.05
        add_card(slide, x, 1.65, 2.85, 2.55, fill=WHITE, border=BORDER)
        add_text(slide, number, x + 0.2, 1.85, 2.45, 0.3, size=10, bold=True, value=BLUE)
        add_text(slide, heading, x + 0.2, 2.25, 2.45, 0.55, size=14, bold=True, line_spacing=1.0)
        add_text(
            slide,
            body,
            x + 0.2,
            2.95,
            2.45,
            1.1,
            size=10.5,
            value=MUTED,
            line_spacing=1.2,
        )

    add_card(slide, 0.7, 4.55, 11.93, 0.6, fill=GREEN_SOFT, border=None)
    add_text(
        slide,
        "Sem conta.   Sem servidor.   Sem envio de dados.",
        0.9,
        4.72,
        11.55,
        0.3,
        size=12.5,
        bold=True,
        value=GREEN,
        align=PP_ALIGN.CENTER,
    )
    add_card(slide, 0.7, 5.45, 11.93, 0.8, fill=NAVY, border=None)
    add_text(
        slide,
        "Windows 10/11   •   Linux x86_64   •   Setup.exe   •   AppImage   •   tar.gz",
        0.9,
        5.67,
        11.55,
        0.35,
        size=11.5,
        bold=True,
        value=WHITE,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Linux: ~/.local/share/learning_helper     Windows: %LOCALAPPDATA%\\learning_helper",
        0.7,
        6.55,
        11.9,
        0.3,
        size=9.5,
        value=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "O app foi feito para ser simples de instalar e discreto. Os dados ficam "
        "em SQLite na máquina do usuário, a instalação padrão não pede root e as "
        "atualizações vêm do GitHub Releases com verificação SHA-256. A "
        "atualização é aplicada quando não há desafio nem lembrete na tela.",
    )

    # ---------------------------------------------------------------- fecho
    slide = prs.slides.add_slide(title_layout)
    add_background(slide, NAVY)
    add_oval(slide, 9.9, -1.5, 4.6, 4.6, NAVY_SOFT)
    add_oval(slide, -2.2, 4.7, 4.8, 4.8, "0B2E36")
    add_text(
        slide,
        "LEARNING HELPER",
        0.85,
        0.9,
        6.0,
        0.3,
        size=10.5,
        bold=True,
        value="5EEAD4",
    )
    set_title(
        slide,
        "Aprender com previsibilidade.\nCrescer com autonomia.",
        x=0.85,
        y=1.55,
        w=8.9,
        h=1.65,
        size=36,
        value=WHITE,
    )
    add_text(
        slide,
        "Pequenos desafios, uma rotina mais leve e dados que ajudam a acompanhar o progresso.",
        0.9,
        3.55,
        7.1,
        0.8,
        size=16,
        value="CBD5E1",
        line_spacing=1.2,
    )
    link = add_text(
        slide,
        "github.com/herossa/learning_helper  ↗",
        0.9,
        5.05,
        6.0,
        0.35,
        size=14,
        bold=True,
        value="5EEAD4",
    )
    link.click_action.hyperlink.address = "https://github.com/herossa/learning_helper"
    add_text(
        slide,
        "Obrigado.",
        0.9,
        6.25,
        4.0,
        0.45,
        size=18,
        bold=True,
        value=WHITE,
    )
    for index, label in enumerate(("7 módulos", "1–5 min", "100% local")):
        x = 9.15 + index * 1.35
        add_card(slide, x, 5.45, 1.2, 0.42, fill=NAVY_SOFT, border="334155", radius=0.5)
        add_text(
            slide,
            label,
            x,
            5.55,
            1.2,
            0.2,
            size=8.5,
            bold=True,
            value="E2E8F0",
            align=PP_ALIGN.CENTER,
        )
    add_notes(
        slide,
        "Para fechar: o Learning Helper é uma forma simples de criar pequenos "
        "momentos de aprendizagem e apoiar a rotina, com previsibilidade e sem "
        "sobrecarga. O código e as releases estão no GitHub.",
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        type=Path,
        default=DEFAULT_OUT,
        help="arquivo .pptx de saída",
    )
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="learning-helper-slides-") as tmp:
        build(args.out, Path(tmp))
    print(f"Apresentação salva em {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
